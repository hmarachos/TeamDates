from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from services.employee_service import EmployeeService


class ExcelService:
    HEADERS = [
        "Наименование предприятия",
        "ФИО",
        "Должность",
        "Подразделение",
        "Дата рождения",
    ]
    HEADER_ALIASES = {
        "company_name": {"наименование предприятия", "предприятие", "организация"},
        "full_name": {"фио", "ф.и.о.", "сотрудник"},
        "position": {"должность"},
        "department": {"подразделение", "отдел"},
        "birth_date": {"дата рождения", "день рождения"},
        "notes": {"комментарий", "примечание", "примечение", "notes"},
    }
    REQUIRED_FIELDS = {"company_name", "full_name", "position", "birth_date"}

    def __init__(self, employee_service: EmployeeService | None = None) -> None:
        self.employee_service = employee_service or EmployeeService()

    def import_employees(self, file: FileStorage) -> dict:
        original_filename = file.filename or ""
        secure_filename(original_filename)
        if Path(original_filename).suffix.lower() != ".xlsx":
            raise ValueError("Поддерживаются только файлы .xlsx.")

        workbook = load_workbook(file, data_only=True)
        sheet = workbook.active
        header_row, header_map = self._find_header_map(sheet)

        imported = 0
        skipped = 0
        errors: list[str] = []
        seen: set[tuple[str, str, date]] = set()

        data_rows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
        for index, row in enumerate(data_rows, start=header_row + 1):
            if self._is_empty_row(row):
                skipped += 1
                continue
            try:
                data = self._row_to_employee(row, header_map)
                key = (
                    data["company_name"].lower(),
                    data["full_name"].lower(),
                    data["birth_date"],
                )
                if key in seen:
                    skipped += 1
                    errors.append(f"Строка {index}: дубликат внутри файла.")
                    continue
                seen.add(key)
                self.employee_service.create(data)
                imported += 1
            except ValueError as exc:
                skipped += 1
                errors.append(f"Строка {index}: {exc}")

        return {"imported": imported, "skipped": skipped, "errors": errors}

    def export_employees(self) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Сотрудники"
        sheet.append(self.HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")

        for employee in self.employee_service.repository.all():
            sheet.append(
                [
                    employee.company_name,
                    employee.full_name,
                    employee.position,
                    employee.department or "",
                    employee.birth_date.strftime("%d.%m.%Y"),
                ]
            )

        widths = [32, 34, 28, 28, 16]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    def _row_to_employee(self, row: tuple, header_map: dict[str, int] | None = None) -> dict:
        header_map = header_map or {
            "company_name": 0,
            "full_name": 1,
            "position": 2,
            "department": 3,
            "birth_date": 4,
        }
        company_name = self._required_text(
            self._cell(row, header_map["company_name"]),
            "Наименование предприятия",
        )
        full_name = self._required_text(self._cell(row, header_map["full_name"]), "ФИО")
        position = self._required_text(self._cell(row, header_map["position"]), "Должность")
        department = self._optional_text(self._cell(row, header_map.get("department")))
        birth_date = self._parse_date(self._cell(row, header_map["birth_date"]))
        notes = self._optional_text(self._cell(row, header_map.get("notes")))
        return {
            "company_name": company_name,
            "full_name": full_name,
            "position": position,
            "department": department,
            "birth_date": birth_date,
            "notes": notes,
        }

    def _find_header_map(self, sheet) -> tuple[int, dict[str, int]]:
        for row_index in range(1, min(sheet.max_row, 20) + 1):
            values = [
                self._normalize_header(sheet.cell(row_index, column_index).value)
                for column_index in range(1, sheet.max_column + 1)
            ]
            header_map: dict[str, int] = {}
            for field, aliases in self.HEADER_ALIASES.items():
                for column_index, value in enumerate(values):
                    if value in aliases:
                        header_map[field] = column_index
                        break
            if self.REQUIRED_FIELDS.issubset(header_map):
                return row_index, header_map
        raise ValueError(
            "Не найдены обязательные колонки: Наименование предприятия, ФИО, "
            "Должность, Дата рождения."
        )

    def _parse_date(self, value) -> date:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
                try:
                    return datetime.strptime(value.strip(), fmt).date()
                except ValueError:
                    continue
        raise ValueError("Некорректный формат даты рождения.")

    def _required_text(self, value, field_name: str) -> str:
        text = self._optional_text(value)
        if not text:
            raise ValueError(f"Поле '{field_name}' обязательно.")
        return text

    def _optional_text(self, value) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    def _cell(self, row: tuple, index: int | None):
        if index is None or index >= len(row):
            return None
        return row[index]

    def _normalize_header(self, value) -> str:
        if value is None:
            return ""
        return " ".join(str(value).strip().lower().replace("ё", "е").split())

    def _is_empty_row(self, row: tuple) -> bool:
        return all(value is None or str(value).strip() == "" for value in row)
