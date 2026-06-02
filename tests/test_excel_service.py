from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

from services.excel_service import ExcelService


def make_file(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(ExcelService.HEADERS)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return FileStorage(stream=stream, filename="Общий список.xlsx")


def test_import_excel(app_ctx):
    service = ExcelService()
    file = make_file(
        [
            ["Компания", "Иванов Иван", "Инженер", "ИТ", "15.07.1985"],
            [None, None, None, None, None],
            ["Компания", "Иванов Иван", "Инженер", "ИТ", "15.07.1985"],
        ]
    )

    result = service.import_employees(file)

    assert result["imported"] == 1
    assert result["skipped"] == 2


def test_import_excel_with_header_on_second_row_and_without_department(app_ctx):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, None, None, None, None])
    sheet.append([
        "Наименование предприятия",
        "ФИО",
        "Должность",
        "Дата рождения",
        "Примечение",
    ])
    sheet.append(["ЦА", "Иванов Иван", "Инженер", "15.07.1985", "важно"])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    file = FileStorage(stream=stream, filename="employees.xlsx")

    result = ExcelService().import_employees(file)

    assert result["imported"] == 1
    assert result["skipped"] == 0


def test_export_excel(app_ctx):
    service = ExcelService()
    service.employee_service.create(
        {
            "company_name": "Компания",
            "full_name": "Иванов Иван",
            "position": "Инженер",
            "department": "ИТ",
            "birth_date": date(1985, 7, 15),
            "notes": None,
        }
    )

    stream = service.export_employees()
    workbook = load_workbook(stream)
    sheet = workbook.active

    assert sheet["A1"].value == "Наименование предприятия"
    assert sheet["B2"].value == "Иванов Иван"
